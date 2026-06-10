import os
# pyrefly: ignore [missing-import]
from docx import Document
from docx.shared import Pt

def create_docx(md_path, docx_path):
    doc = Document()
    
    try:
        with open(md_path, 'r', encoding='utf-8') as f:
            lines = f.readlines()
            
        for line in lines:
            line = line.strip()
            if not line or line == '---':
                continue
                
            if line.startswith('# '):
                p = doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                p = doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                p = doc.add_heading(line[4:], level=3)
            elif line.startswith('- '):
                # Handle bold inside lists
                p = doc.add_paragraph(style='List Bullet')
                parts = line[2:].split('**')
                for i, part in enumerate(parts):
                    run = p.add_run(part)
                    if i % 2 != 0:
                        run.bold = True
            elif line.startswith('* '):
                p = doc.add_paragraph(style='List Bullet')
                parts = line[2:].split('**')
                for i, part in enumerate(parts):
                    run = p.add_run(part)
                    if i % 2 != 0:
                        run.bold = True
            elif line[0].isdigit() and line[1:3] == '. ':
                p = doc.add_paragraph(style='List Number')
                parts = line[3:].split('**')
                for i, part in enumerate(parts):
                    run = p.add_run(part)
                    if i % 2 != 0:
                        run.bold = True
            else:
                p = doc.add_paragraph()
                parts = line.split('**')
                for i, part in enumerate(parts):
                    run = p.add_run(part)
                    if i % 2 != 0:
                        run.bold = True
                        
        doc.save(docx_path)
        print("Success")
    except Exception as e:
        print(f"Error: {e}")

if __name__ == '__main__':
    md_file = "project_report.md" # Changed from hardcoded absolute path
    docx_file = "Project_Report.docx"
    if os.path.exists(md_file):
        create_docx(md_file, docx_file)
    else:
        print(f"Error: {md_file} not found.")
